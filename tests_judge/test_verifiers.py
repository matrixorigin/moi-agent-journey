"""Exercise the document checkers against small, editable Office fixtures."""
from pathlib import Path
from io import BytesIO
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from zipfile import ZipFile

import fitz
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from judge import report_title, statement, summary_sheet, table_chart
from judge.verify import adapt_skillsbench_states, check, preflight


class DocumentVerifiersTest(unittest.TestCase):
    def test_office_preflight_parses_documents_instead_of_only_checking_zip(self):
        with tempfile.TemporaryDirectory() as directory:
            outputs=Path(directory)
            Document().save(outputs/'valid.docx')
            with ZipFile(outputs/'broken.docx','w') as archive:
                archive.writestr('[Content_Types].xml','<broken')
            checks={item['rule']:item for item in preflight(outputs)}
            self.assertTrue(checks['文件可解析：valid.docx']['passed'])
            self.assertFalse(checks['文件可解析：broken.docx']['passed'])

    def test_report_title_comes_from_source_and_honors_inherited_formatting(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); inputs = root/'inputs'; outputs = root/'outputs'; rendered = root/'rendered/report.docx'
            inputs.mkdir(); outputs.mkdir(); rendered.mkdir(parents=True)
            source = Document(); source.add_paragraph('Sample report'); source.save(inputs/'公共书房运营规划书.docx')
            doc = Document(); normal = doc.styles['Normal']; normal.font.name = 'SimSun'; normal.font.size = Pt(16); normal.font.bold = True
            title = doc.add_paragraph('Sample report'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title.paragraph_format.line_spacing = 1.0
            abstract_style = doc.styles.add_style('AbstractForTest', WD_STYLE_TYPE.PARAGRAPH)
            abstract_style.font.name = 'SimHei'; abstract_style.font.size = Pt(14)
            doc.add_paragraph('摘要', style=abstract_style); doc.save(outputs/'report.docx')
            self.assertTrue(all(item['passed'] for item in report_title.verify(inputs, outputs, check)))
            pdf = fitz.open(); page = pdf.new_page(); page.insert_text((40,40), 'Sample report')
            pdf.save(str(rendered/'report.pdf')); pdf.close()
            self.assertTrue(report_title.verify_rendered(inputs, root/'rendered', check)['passed'])
            doc = Document(outputs/'report.docx'); doc.paragraphs[0].runs[0].font.bold = False
            doc.save(outputs/'report.docx')
            self.assertFalse(report_title.verify(inputs, outputs, check)[0]['passed'])
            doc = Document(outputs/'report.docx'); doc.paragraphs[0].text = 'Wrong report'
            doc.save(outputs/'report.docx')
            self.assertFalse(report_title.verify(inputs, outputs, check)[0]['passed'])

    def test_presentation_requires_all_source_data_and_matching_embedded_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); inputs = root/'inputs'; outputs = root/'outputs'
            inputs.mkdir(); outputs.mkdir()
            headings = ['年份', '响应率（%）', '协同指数（0–100）']
            rows = [[2023, 10.0, 30.0], [2024, 20.0, 40.0]]
            workbook = Workbook(); sheet = workbook.active; sheet.title = '关联不同数据'
            for row in [headings, *rows]: sheet.append(row)
            workbook.save(inputs/'PPT相关不同数据表.xlsx'); workbook.close()
            source = Presentation(); first = source.slides.add_slide(source.slide_layouts[6])
            first.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = 'Existing slide'
            source.save(inputs/'表格图表关联数据.pptx')
            output = Presentation(inputs/'表格图表关联数据.pptx')
            table_slide = output.slides.add_slide(output.slide_layouts[6])
            table = table_slide.shapes.add_table(3, 3, Inches(1), Inches(1), Inches(6), Inches(2)).table
            for r, values in enumerate([headings, *rows]):
                for c, value in enumerate(values): table.cell(r, c).text = str(value)
            chart_slide = output.slides.add_slide(output.slide_layouts[6]); data = CategoryChartData()
            data.categories = [row[0] for row in rows]
            data.add_series('响应率', [row[1] for row in rows]); data.add_series('协同指数', [row[2] for row in rows])
            chart_slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(1), Inches(1), Inches(6), Inches(4), data)
            target = outputs/'report.pptx'; output.save(target)
            checks, _ = table_chart.verify(inputs, outputs, check)
            self.assertTrue(all(item['passed'] for item in checks), checks)
            changed = Presentation(target); changed.slides[-2].shapes[0].table.cell(1, 1).text = '11'
            changed.save(target)
            checks, _ = table_chart.verify(inputs, outputs, check)
            self.assertFalse(checks[2]['passed'])
            output.save(target)
            changed = Presentation(target); chart = changed.slides[-1].shapes[0].chart
            cache = chart.series[0]._element.xpath('./c:val/c:numRef/c:numCache/c:pt/c:v')
            cache[0].text = '999'; changed.save(target)
            checks, _ = table_chart.verify(inputs, outputs, check)
            self.assertFalse(checks[3]['passed'])
            output.save(target)
            changed = Presentation(target); chart = changed.slides[-1].shapes[0].chart
            part = chart.part.chart_workbook.xlsx_part
            embedded = load_workbook(BytesIO(part.blob)); embedded.active['B2'] = 999
            buffer = BytesIO(); embedded.save(buffer); embedded.close(); part._blob = buffer.getvalue()
            changed.save(target)
            checks, _ = table_chart.verify(inputs, outputs, check)
            self.assertFalse(checks[3]['passed'])
            output.save(target)
            changed = Presentation(target); chart = changed.slides[-1].shapes[0].chart
            series = chart.series[-1]._element; series.getparent().remove(series); changed.save(target)
            checks, _ = table_chart.verify(inputs, outputs, check)
            self.assertFalse(checks[3]['passed'])
            output.save(target)
            changed = Presentation(target); changed.slides[0].shapes[0].text = 'Changed original'
            changed.save(target)
            checks, _ = table_chart.verify(inputs, outputs, check)
            self.assertFalse(checks[1]['passed'])

    def test_statement_requires_source_and_original_records_with_correct_amounts(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); inputs = root/'inputs'; outputs = root/'outputs'
            inputs.mkdir(); outputs.mkdir()
            pdf = fitz.open(); page = pdf.new_page()
            page.insert_text((40, 40), '\n'.join([
                '01 Dec. 2025 to 31 Dec. 2025',
                '01 Dec. Balance forward $100.00',
                '01 Dec. Insurance 20.00 80.00',
                '02 Dec. Sales 50.00 130.00',
                '02 Dec. Closing totals 20.00 50.00 130.00']))
            pdf.save(str(inputs/'Cedarline_Statement_2025DEC_reordered.pdf')); pdf.close()
            header = ['Date', 'Description', 'Withdrawals', 'Deposits', 'Balance']
            original = ['Mar 22', 'Old transaction', '', '$10.00', '$70.00']
            records = [original, ['01 Dec', 'Insurance', '-$20.00', '', '$80.00'],
                       ['02 Dec', 'Sales', '', '$50.00', '$130.00']]
            def write(path, values):
                doc = Document(); table = doc.add_table(rows=1, cols=5)
                for cell, value in zip(table.rows[0].cells, header): cell.text = value
                for values_row in values:
                    row = table.add_row(); row.height = Pt(20); row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
                    for cell, value in zip(row.cells, values_row): cell.text = value
                doc.save(str(path))
            write(inputs/'public_ready_statement.docx', [original])
            output = outputs/'statement.docx'
            write(output, records)
            self.assertTrue(all(item['passed'] for item in statement.verify(inputs, outputs, check)[0]))
            opening = ['01 Dec', 'Balance forward', '', '', '$100.00']
            write(output, [opening] + records)
            self.assertTrue(all(item['passed'] for item in statement.verify(inputs, outputs, check)[0]))
            for invalid in (["01 Dec", "Balance forward", "", "", "$99.00"],
                            ["02 Dec", "Balance forward", "", "", "$100.00"],
                            ["01 Dec 2024", "Balance forward", "", "", "$100.00"],
                            ["01 Dec", "Balance forward", "$1.00", "", "$100.00"]):
                write(output, [invalid] + records)
                self.assertFalse(statement.verify(inputs, outputs, check)[0][1]['passed'])
            write(output, [opening, opening] + records)
            self.assertFalse(statement.verify(inputs, outputs, check)[0][1]['passed'])
            write(output, records)
            visible = ' '.join(value for row in records for value in row if value)
            self.assertEqual(statement.rendered_cells_visible(Document(str(output)), visible), [])
            self.assertTrue(statement.rendered_cells_visible(Document(str(output)), visible.replace('Insurance', '')))
            write(output, records[1:])
            self.assertFalse(statement.verify(inputs, outputs, check)[0][1]['passed'])
            changed = [row[:] for row in records]; changed[1][2] = '-$21.00'
            write(output, changed)
            self.assertFalse(statement.verify(inputs, outputs, check)[0][1]['passed'])
            changed = [row[:] for row in records]; changed[1][0] = '01 Dec 2024'
            write(output, changed)
            self.assertFalse(statement.verify(inputs, outputs, check)[0][1]['passed'])
            write(output, records + [records[1]])
            self.assertFalse(statement.verify(inputs, outputs, check)[0][1]['passed'])

    def test_summary_uses_source_labels_and_detects_wrong_values_or_changes(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); inputs = root/'inputs'; outputs = root/'outputs'
            inputs.mkdir(); outputs.mkdir()
            workbook = Workbook(); summary = workbook.active; summary.title = '汇总'
            fixtures = [
                ('Sheet1', [('A1', '总重量合计'), ('B1', 19), ('B3', '砼体积(m³)'), ('B4', 0.5)]),
                ('Sheet2', [('A1', 17), ('B1', '总重量合计'), ('C3', '砼体积(m³)'), ('C4', 0.25)]),
                ('Sheet3', [('B1', '总重量合计'), ('B2', 12), ('C3', '砼体积(m³)'), ('C4', 0.4)])]
            for row, (name, cells) in enumerate(fixtures, 3):
                summary.cell(row, 3, name); summary.cell(row, 8, row * 2)
                sheet = workbook.create_sheet(name)
                for coordinate, value in cells: sheet[coordinate] = value
            source = inputs/'子表数据提取到汇总表.xlsx'; workbook.save(source)
            summary['F3'] = '=Sheet1!B4'; summary['G3'] = '=Sheet1!B1'
            summary['F4'] = 0.25; summary['G4'] = 17
            summary['F5'] = '=Sheet3!C4'; summary['G5'] = '=Sheet3!B2'
            target = outputs/'summary.xlsx'; workbook.save(target); workbook.close()
            checks, _ = summary_sheet.verify(inputs, outputs, check)
            self.assertTrue(all(item['passed'] for item in checks), checks)
            with patch('judge.summary_sheet.subprocess.run', return_value=SimpleNamespace(
                    returncode=-11, stdout='', stderr='process crashed')):
                with self.assertRaisesRegex(RuntimeError, 'summary_recalculation_process_failed'):
                    summary_sheet.verify(inputs, outputs, check)
            with patch('judge.summary_sheet.subprocess.run', return_value=SimpleNamespace(
                    returncode=0, stdout='', stderr='source file could not be loaded')):
                with self.assertRaisesRegex(RuntimeError, 'summary_recalculation_no_output'):
                    summary_sheet.verify(inputs, outputs, check)
            changed = load_workbook(target); changed['汇总']['F3'] = 0.75; changed.save(target); changed.close()
            self.assertFalse(summary_sheet.verify(inputs, outputs, check)[0][1]['passed'])
            changed = load_workbook(target); changed['汇总']['F3'] = 0.5; changed['Sheet2']['A1'] = 18
            changed.save(target); changed.close()
            checks, _ = summary_sheet.verify(inputs, outputs, check)
            self.assertTrue(checks[1]['passed'], checks)
            self.assertFalse(checks[2]['passed'])

    def test_population_adapter_requires_source_evidence(self):
        from types import SimpleNamespace
        with tempfile.TemporaryDirectory() as directory:
            inputs = Path(directory); module = SimpleNamespace(VALID_STATES={'Victoria'})
            pdf = fitz.open(); page = pdf.new_page(); page.insert_text((40, 40), 'Other Territories')
            pdf.save(str(inputs/'population.pdf')); pdf.close()
            adapt_skillsbench_states({'id': 'J1-08'}, module, inputs)
            self.assertIn('Other Territories', module.VALID_STATES)
            (inputs/'population.pdf').unlink()
            pdf = fitz.open(); page = pdf.new_page(); page.insert_text((40, 40), 'Victoria')
            pdf.save(str(inputs/'population.pdf')); pdf.close()
            with self.assertRaisesRegex(RuntimeError, 'population_source_missing_other_territories'):
                adapt_skillsbench_states({'id': 'J1-08'}, module, inputs)


if __name__ == '__main__':
    unittest.main()
