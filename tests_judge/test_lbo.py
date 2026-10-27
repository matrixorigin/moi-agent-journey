"""Exercise financial checks without constraining unrelated valid formula repairs."""
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from judge.verify import lbo


class LBOTest(unittest.TestCase):
    def test_valid_formula_repair_and_material_regressions(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); inputs = root/'inputs'; outputs = root/'outputs'
            inputs.mkdir(); outputs.mkdir()
            source = Workbook(); sheet = source.active; sheet.title = 'LBO'
            financials = source.create_sheet('Financials'); source.create_sheet('Operating Cases')
            for col in range(7, 15):
                c = get_column_letter(col)
                for row, value in [(57, 110), (64, 100), (67, -10), (68, -20), (75, 2),
                                   (136, 1), (137, 2), (145, 3), (152, 4)]:
                    sheet[f'{c}{row}'] = value
                sheet[f'{c}66'] = f'={c}64+{c}75'
                sheet[f'{c}69'] = f'={c}66+{c}67+{c}68'
                sheet[f'{c}175'] = f'={c}136+{c}137+{c}145+{c}152'
            for col in range(10, 15):
                c = get_column_letter(col)
                for row, value in [(12, 100), (15, 10), (16, 5), (17, 4)]:
                    financials[f'{c}{row}'] = value
                financials[f'{c}18'] = f'={c}12-{c}15-{c}16-{c}17'
            sheet['G50'] = '=1+1'
            source.calculation.calcMode = 'auto'; source.calculation.fullCalcOnLoad = True
            source.save(inputs/'09_01_input.xlsx'); source.close()
            output_path = outputs/'09_01_repaired.xlsx'
            workbook = load_workbook(inputs/'09_01_input.xlsx')
            workbook['LBO']['G50'] = '=SUM(1,1)'
            workbook.save(output_path); workbook.close()

            def checked():
                checks, _ = lbo(None, inputs, outputs)
                return {item['rule']: item for item in checks}

            checks = checked()
            self.assertTrue(all(item['passed'] for item in checks.values()), checks)
            preservation = '原非目标数值未被静态改动、公式仍可编辑且合并结构保留'
            workbook = load_workbook(output_path); workbook['LBO']['G50'] = 2
            workbook.save(output_path); workbook.close()
            self.assertFalse(checked()[preservation]['passed'])
            workbook = load_workbook(output_path); workbook['LBO']['G50'] = '=SUM(1,1)'
            workbook['LBO']['G64'] = 999; workbook.save(output_path); workbook.close()
            self.assertFalse(checked()[preservation]['passed'])
            workbook = load_workbook(output_path); workbook['LBO']['G64'] = 100
            workbook['LBO']['G66'] = '=G57+G75'; workbook.save(output_path); workbook.close()
            self.assertFalse(checked()['LBO!G66 的业务关系和原生公式']['passed'])


if __name__ == '__main__':
    unittest.main()
