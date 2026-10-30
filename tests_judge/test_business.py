import json
import shutil
import tempfile
import unittest
from pathlib import Path

from judge.business import extract, verify


def check(rule,passed,detail=''):
    return {'rule':rule,'passed':bool(passed),'detail':detail}


class BusinessJudgeTests(unittest.TestCase):
    def case(self,n,files):
        return {'source':{'task':f'T{n:03d}'},'expected_files':files}

    def test_wide_table_checks_numbers_not_just_presence(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);(root/'long.csv').write_text('product,month,amount\nA,1,100\nA,2,120\nA,3,150\nB,1,200\nB,2,200\nB,3,180\n')
            checks,_=verify(self.case(28,['long.csv']),root,root,check)
            self.assertTrue(all(x['passed'] for x in checks))
            (root/'long.csv').write_text((root/'long.csv').read_text().replace('150','151'))
            checks,_=verify(self.case(28,['long.csv']),root,root,check)
            self.assertFalse(all(x['passed'] for x in checks))

    def test_copied_files_do_not_pass_move_task(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);inputs=root/'inputs';out=root/'outputs'
            source=inputs/'samples/T091_downloads';source.mkdir(parents=True)
            for n in range(3):(source/f'{n}.txt').write_text(str(n))
            shutil.copytree(source,out/'T091_downloads')
            checks,_=verify(self.case(91,[]),inputs,out,check)
            self.assertFalse(all(x['passed'] for x in checks))
            target=out/'T091_downloads/classified';target.mkdir()
            for p in list((out/'T091_downloads').glob('*.txt')):p.rename(target/p.name)
            checks,_=verify(self.case(91,[]),inputs,out,check)
            self.assertTrue(all(x['passed'] for x in checks))
            (target/'0.txt').write_text('changed')
            checks,_=verify(self.case(91,[]),inputs,out,check)
            self.assertFalse(all(x['passed'] for x in checks))

    def test_pdf_and_office_extraction_uses_real_binary_content(self):
        import fitz
        from docx import Document
        from pptx import Presentation
        from pptx.util import Inches
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            doc=Document();doc.add_paragraph('delivered document');doc.save(root/'manual.docx')
            slides=Presentation()
            for n in range(10):
                slide=slides.slides.add_slide(slides.slide_layouts[6])
                slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(2)).text='slide '+str(n)
                slide.notes_slide.notes_text_frame.text='speaker notes '+str(n)
            slides.save(root/'slides.pptx')
            book=Workbook();book.active['A1']='=100*1.13';book.save(root/'quotes.xlsx')
            pdf=fitz.open();p=pdf.new_page();p.insert_text((72,72),'source page');pdf.save(root/'summary.pdf');pdf.close()
            data=extract(root)
            self.assertIn('delivered document',data['manual.docx'])
            self.assertIn('speaker notes 9',data['slides.pptx'])
            self.assertIn('=100*1.13',data['quotes.xlsx'])
            self.assertIn('source page',data['summary.pdf'])
            checks,_=verify(self.case(66,['slides.pptx']),root,root,check)
            self.assertTrue(all(x['passed'] for x in checks))
            checks,_=verify(self.case(63,['summary.pdf']),root,root,check)
            self.assertFalse(all(x['passed'] for x in checks))

    def test_image_task_requires_three_real_images(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);source=root/'inputs/samples/T094_images';source.mkdir(parents=True)
            out=root/'outputs';out.mkdir()
            for n,size in enumerate([(1200,800),(800,1200),(1200,1200)]):
                Image.new('RGB',(size[0]*2,size[1]*2)).save(source/f'{n}.png')
                Image.new('RGB',size).save(out/f'{n}.png')
            checks,_=verify(self.case(94,[]),root/'inputs',out,check)
            self.assertTrue(all(x['passed'] for x in checks))
            Image.new('RGB',(1300,1300)).save(out/'0.png')
            checks,_=verify(self.case(94,[]),root/'inputs',out,check)
            self.assertFalse(all(x['passed'] for x in checks))
