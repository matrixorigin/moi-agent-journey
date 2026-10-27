"""Verify editable PowerPoint data and its embedded workbook against source Excel."""
from hashlib import sha256
from io import BytesIO
import math
import re


def label(value):
    value = str(value or '').strip()
    value = re.sub(r'[（(]\s*(?:%|0\s*[–-]\s*100)\s*[）)]$', '', value)
    return ''.join(value.split()).casefold()


def number(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        number_value = float(str(value).strip().replace(',', ''))
        return number_value if math.isfinite(number_value) else None
    except ValueError:
        return None


def equal_number(actual, expected):
    actual_number, expected_number = number(actual), number(expected)
    return actual_number is not None and expected_number is not None and math.isclose(
        actual_number, expected_number, rel_tol=1e-9, abs_tol=1e-8)


def shape_content(shape):
    result = [int(shape.shape_type), shape.left, shape.top, shape.width, shape.height, shape.rotation]
    if shape.has_text_frame:
        result.append(('text', shape.text))
    if shape.has_table:
        result.append(('table', tuple(tuple(cell.text for cell in row.cells) for row in shape.table.rows)))
    if shape.has_chart:
        chart = shape.chart
        result.append(('chart', int(chart.chart_type),
            tuple(str(category.label) for category in chart.plots[0].categories),
            tuple((series.name, tuple(series.values)) for series in chart.series)))
    if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
        result.append(('image', sha256(shape.image.blob).hexdigest()))
    return tuple(result)


def slide_content(slide):
    return tuple(shape_content(shape) for shape in slide.shapes)


def cells_from_reference(workbook, formula):
    from openpyxl.utils.cell import range_boundaries
    if not formula or '!' not in formula:
        raise ValueError('missing_embedded_data_reference')
    sheet_name, cell_range = formula.rsplit('!', 1)
    sheet_name = sheet_name.strip("'").replace("''", "'")
    if '[' in sheet_name or ']' in sheet_name or sheet_name not in workbook.sheetnames:
        raise ValueError('external_or_missing_chart_sheet')
    min_col, min_row, max_col, max_row = range_boundaries(cell_range.replace('$', ''))
    if min_col != max_col and min_row != max_row:
        raise ValueError('chart_reference_is_not_one_dimensional')
    return [cell.value for row in workbook[sheet_name].iter_rows(
        min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col) for cell in row]


def table_errors(table, headings, rows):
    if len(table.rows) != len(rows) + 1 or len(table.columns) != len(headings):
        return ['table_dimensions']
    errors = []
    for col, heading in enumerate(headings):
        if label(table.cell(0, col).text) != label(heading):
            errors.append('header:' + str(col))
    for row_index, source_row in enumerate(rows, 1):
        for col, expected in enumerate(source_row):
            if not equal_number(table.cell(row_index, col).text, expected):
                errors.append(f'cell:{row_index}:{col}')
    return errors


def chart_errors(chart, headings, rows):
    from openpyxl import load_workbook
    from pptx.enum.chart import XL_CHART_TYPE
    line_types = {XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS}
    if chart.chart_type not in line_types or len(chart.plots) != 1:
        return ['not_a_single_editable_line_chart']
    errors = []
    years = [row[0] for row in rows]
    categories = [category.label for category in chart.plots[0].categories]
    if len(categories) != len(years) or any(not equal_number(a, b) for a, b in zip(categories, years)):
        errors.append('displayed_years')
    expected = {label(heading): [row[index] for row in rows]
                for index, heading in enumerate(headings[1:], 1)}
    actual_names = [label(series.name) for series in chart.series]
    if len(actual_names) != len(set(actual_names)) or set(actual_names) != set(expected):
        errors.append('displayed_series')
    try:
        embedded = load_workbook(BytesIO(chart.part.chart_workbook.xlsx_part.blob), data_only=True)
    except (AttributeError, KeyError, ValueError, OSError):
        return errors + ['embedded_workbook_missing_or_invalid']
    try:
        for series in chart.series:
            name = label(series.name)
            if name not in expected:
                continue
            target = expected[name]
            values = list(series.values)
            if len(values) != len(target) or any(not equal_number(a, b) for a, b in zip(values, target)):
                errors.append('displayed_values:' + name)
            try:
                value_refs = series._element.xpath('./c:val/c:numRef/c:f')
                category_refs = series._element.xpath('./c:cat/c:numRef/c:f|./c:cat/c:strRef/c:f')
                name_refs = series._element.xpath('./c:tx/c:strRef/c:f')
                if len(value_refs) != 1 or len(category_refs) != 1 or len(name_refs) != 1:
                    raise ValueError('chart_data_references_missing')
                workbook_values = cells_from_reference(embedded, value_refs[0].text)
                workbook_years = cells_from_reference(embedded, category_refs[0].text)
                workbook_names = cells_from_reference(embedded, name_refs[0].text)
                if len(workbook_values) != len(target) or any(
                        not equal_number(a, b) for a, b in zip(workbook_values, target)):
                    errors.append('embedded_values:' + name)
                if len(workbook_years) != len(years) or any(
                        not equal_number(a, b) for a, b in zip(workbook_years, years)):
                    errors.append('embedded_years:' + name)
                if len(workbook_names) != 1 or label(workbook_names[0]) != name:
                    errors.append('embedded_name:' + name)
            except (ValueError, KeyError, TypeError, IndexError):
                errors.append('embedded_references:' + name)
    finally:
        embedded.close()
    return errors


def verify(inputs, outputs, check):
    from openpyxl import load_workbook
    from pptx import Presentation
    workbook = load_workbook(inputs/'PPT相关不同数据表.xlsx', data_only=True)
    sheet = workbook['关联不同数据']
    values = [list(row) for row in sheet.values]
    workbook.close()
    if not values or not values[0] or any(len(row) != len(values[0]) for row in values):
        raise ValueError('source_data_table_invalid')
    headings, rows = values[0], values[1:]
    if (not rows or any(not label(h) for h in headings) or
            len({label(h) for h in headings}) != len(headings) or
            any(any(number(value) is None for value in row) for row in rows)):
        raise ValueError('source_data_table_invalid')
    original = Presentation(inputs/'表格图表关联数据.pptx')
    output = Presentation(outputs/'report.pptx')
    expected_count = len(original.slides) + 2
    checks = [check('原稿后依次新增表格页和图表页', len(output.slides) == expected_count,
                    str({'original': len(original.slides), 'output': len(output.slides)}))]
    if len(output.slides) != expected_count:
        return checks, {'checks': checks}
    changed = [index + 1 for index, slide in enumerate(original.slides)
               if slide_content(slide) != slide_content(output.slides[index])]
    checks.append(check('原有幻灯片内容和布局保留', not changed, str({'changed_slides': changed})))
    tables = [shape.table for shape in output.slides[-2].shapes if shape.has_table]
    table_problems = table_errors(tables[0], headings, rows) if len(tables) == 1 else ['editable_table_count']
    checks.append(check('原生表格包含源数据的全部年份和指标', not table_problems, str(table_problems[:20])))
    charts = [shape.chart for shape in output.slides[-1].shapes if shape.has_chart]
    chart_problems = chart_errors(charts[0], headings, rows) if len(charts) == 1 else ['editable_chart_count']
    checks.append(check('原生折线图与嵌入工作簿均匹配全部源数据', not chart_problems,
                        str(chart_problems[:20])))
    return checks, {'checks': checks, 'source_rows': len(rows), 'source_series': len(headings)-1}
