"""Validate AP control packages against source invoices, formulas and drafts."""
from collections import Counter
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from email.utils import getaddresses
import json
import mailbox
import re
import subprocess
import tempfile
from pathlib import Path
from zipfile import ZipFile

from .office_process import require_success


HEADERS = {
    'invoice': {'invoice', 'invoiceno', 'invoiceid', 'linkedinvoice'},
    'vendor': {'vendor', 'vendoraswritten'},
    'source': {'sourcemessageid', 'messageid', 'messageidref', 'sourceemailid'},
    'sender': {'sender', 'from'},
    'facts': {'factextracted'},
    'source_amount': {'sourceamount', 'amountasstated', 'amountasinvoiced', 'emailamountusd'},
    'source_due': {'sourcedue', 'duedateasstated', 'duedateasinvoiced', 'duedatestated'},
    'correction': {'correctionamount', 'correctionamountusd'},
    'amount': {'effectiveamount', 'effectiveamountusd', 'amount'},
    'block': {'blockstatus', 'blockedvendor'},
    'decision': {'disposition'},
    'review': {'reviewflag', 'notes'},
    'due': {'adjusteddue', 'adjustedduedate', 'duedateadjusted', 'duedate'},
    'recipient': {'recipient', 'to'},
    'delivery': {'deliverystatus'},
}


def normalized(value):
    return re.sub(r'[^a-z0-9]+', '', str(value or '').casefold())


def amount(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        result = Decimal(str(value).replace(',', '').replace('$', '').strip())
        return result.quantize(Decimal('0.01')) if result.is_finite() else None
    except (InvalidOperation, ValueError):
        return None


def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    numeric = re.fullmatch(r'(\d{4})\s*[-/.年]\s*(\d{1,2})\s*[-/.月]\s*(\d{1,2})日?', text)
    if numeric:
        try:
            return date(*(int(part) for part in numeric.groups()))
        except ValueError:
            return None
    for pattern in ('%B %d, %Y', '%b %d, %Y', '%d %B %Y', '%d %b %Y'):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def next_workday(value):
    while value.weekday() >= 5:
        value += timedelta(days=1)
    return value


def find_header(workbook, required, sheet_name=None, optional=frozenset()):
    matches = []
    for sheet in workbook:
        if sheet_name and sheet.title != sheet_name:
            continue
        for row in sheet.iter_rows():
            columns = {}
            selected = {}
            ambiguous = set()
            for cell in row:
                value = normalized(cell.value)
                for field in required | optional:
                    if value in HEADERS[field]:
                        if field in columns:
                            if field == 'review' and {value, selected[field]} == {'reviewflag', 'notes'}:
                                if value == 'notes':
                                    continue
                            else:
                                ambiguous.add(field)
                                continue
                        columns[field] = cell.column
                        selected[field] = value
            if required <= set(columns):
                if ambiguous:
                    raise ValueError('ambiguous_header:' + sheet.title + ':' + ','.join(sorted(ambiguous)))
                matches.append((sheet.title, row[0].row, columns))
    if len(matches) != 1:
        raise ValueError('unsupported_or_ambiguous_workbook_layout:' + ','.join(sorted(required)))
    return matches[0]


def invoice_rows(sheet, header_row, invoice_column, invoice_ids, vendor_column):
    rows = {}
    duplicates = []
    extras = []
    for row in range(header_row + 1, sheet.max_row + 1):
        value = str(sheet.cell(row, invoice_column).value or '').strip()
        if value.casefold().startswith(('total ', 'logged total')):
            break
        if value in invoice_ids:
            if value in rows:
                duplicates.append(value)
            rows[value] = row
        elif value and (re.fullmatch(r'[A-Z]{2,8}-[A-Z0-9-]{3,}', value) or
                        sheet.cell(row, vendor_column).value is not None):
            extras.append(value)
    return rows, duplicates, extras


def recalculate(source, destination):
    destination.mkdir(parents=True, exist_ok=True)
    profile = destination.parent / (destination.name + '-profile')
    command = ['libreoffice', '-env:UserInstallation=' + profile.as_uri(), '--headless',
               '--convert-to', 'xlsx', '--outdir', str(destination), str(source)]
    result = subprocess.run(command, capture_output=True, text=True, timeout=120)
    output = destination/source.name
    require_success(result, 'ap_workbook_recalculation_process_failed')
    if not output.is_file():
        raise RuntimeError('ap_workbook_recalculation_failed:' + (result.stdout + result.stderr)[-500:])
    return output


def cell_value(workbook, layout, row, field):
    sheet, _, columns = layout
    return workbook[sheet].cell(row, columns[field]).value


def status_matches(invoice, block, disposition, review, due):
    blocked = normalized(block)
    disposition = normalized(disposition)
    review = str(review or '').strip().casefold()
    if invoice.blocked:
        adjusted = next_workday(invoice.source_due)
        return (blocked in {'blocked', 'activeblock', 'yes', 'true'} and disposition in {'reject', 'rejected'}
                and 'review - over threshold' not in review
                and (not review or 'reject' in review or review in {'n/a', 'na', 'not required', '-'})
                and (due is None or as_date(due) in {invoice.source_due, adjusted}))
    expected_review = invoice.review_flag.casefold()
    return (blocked in {'clear', 'no', 'false', 'notblocked'} and disposition in {'log', 'logged', 'post', 'posted', 'pending'}
            and (expected_review in review if expected_review else 'review - over threshold' not in review)
            and as_date(due) == invoice.adjusted_due)


def summary_metric(workbook, category):
    sheet = workbook['Control_Summary']
    for row in sheet.iter_rows():
        title = str(row[0].value or '').strip().casefold()
        if category == 'reviewed' and (title == 'source invoices' or 'invoices reviewed' in title):
            return row[1]
        if category == 'posted' and (title in {'logged invoices', 'invoices cleared for posting'} or title.startswith('posted to pending_invoices')):
            return row[1]
        if category == 'rejected' and (title in {'rejected invoices', 'invoices rejected'} or title.startswith('rejected (active block)')):
            return row[1]
        if category == 'review' and (title == 'review invoices' or title.startswith(('review flags raised', 'invoices flagged review'))):
            return row[1]
        if category == 'weekend' and (title == 'weekend-adjusted invoices' or title.startswith(('weekend due dates rolled', 'weekend adjustments carried into the pending list'))):
            return row[1]
        if category == 'drafts' and (title in {'communication drafts', 'offline drafts produced'} or title.startswith('offline drafts prepared')):
            return row[1]
    raise ValueError('unsupported_summary_metric:' + category)


def threshold_cell(formulas, values):
    candidates = []
    for sheet in formulas:
        if not normalized(sheet.title).startswith('policy'):
            continue
        for row in sheet.iter_rows():
            content = ' '.join(str(cell.value or '') for cell in row).casefold()
            if not any(token in content for token in ('vp finance', 'msg_vp_override', 'review_threshold')):
                continue
            if 'superseded' in content:
                continue
            for cell in row:
                current = values[sheet.title][cell.coordinate].value
                if isinstance(current, (int, float)) and not isinstance(current, bool):
                    candidates.append(cell)
    if len(candidates) != 1:
        raise ValueError('unsupported_or_ambiguous_threshold_location')
    return candidates[0]


def pending_total(workbook):
    sheet = workbook['Pending_Invoices']
    _, _, fields = find_header(workbook, {'invoice', 'vendor', 'amount', 'due', 'review'}, 'Pending_Invoices')
    candidates = []
    for row in sheet.iter_rows():
        titles = [str(cell.value or '').strip().casefold() for cell in row]
        if any(title.startswith(('logged total', 'total -')) or title in {'total', 'total pending'} for title in titles):
            candidates.extend(cell for cell in row[1:] if cell.data_type == 'f' and cell.column == fields['amount'])
    if len(candidates) != 1:
        raise ValueError('unsupported_or_ambiguous_pending_total')
    return candidates[0]


def body(message):
    chunks = []
    for part in message.walk():
        if part.is_multipart() or part.get_content_type() != 'text/plain':
            continue
        raw = part.get_payload(decode=True)
        if raw is None:
            chunks.append(str(part.get_payload()))
        else:
            chunks.append(raw.decode(part.get_content_charset() or 'utf-8', errors='strict'))
    return '\n'.join(chunks)


def document_items(document):
    return [p.text for p in document.paragraphs] + [
        ' | '.join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]


def document_table_conflicts(document, expected):
    """Check stated invoice fields in recognizable tables without requiring one."""
    problems = []
    for table in document.tables:
        if not table.rows:
            continue
        headers = [normalized(cell.text) for cell in table.rows[0].cells]
        if 'invoice' not in headers:
            continue
        invoice_column = headers.index('invoice')
        for row in table.rows[1:]:
            if invoice_column >= len(row.cells):
                continue
            invoice = expected.get(row.cells[invoice_column].text.strip())
            if invoice is None:
                continue
            for index, heading in enumerate(headers):
                value = row.cells[index].text.strip()
                if not value:
                    continue
                if heading == 'vendor' and value.casefold() != invoice.vendor.casefold():
                    problems.append(invoice.invoice_id + ':vendor')
                if heading in {'amount', 'effective', 'effectiveamount', 'asinvoiced', 'sourceamount'}:
                    stated = amount(value)
                    wanted = invoice.source_amount if heading in {'asinvoiced', 'sourceamount'} else invoice.effective_amount
                    if stated is not None and stated != wanted:
                        problems.append(invoice.invoice_id + ':amount')
                if heading == 'disposition':
                    disposition = normalized(value)
                    if (invoice.blocked and disposition in {'log', 'logged', 'post', 'posted', 'pending'}) or (
                            not invoice.blocked and disposition in {'reject', 'rejected', 'block', 'blocked'}):
                        problems.append(invoice.invoice_id + ':disposition')
                if heading in {'duedate', 'adjusteddue'}:
                    stated = as_date(value)
                    wanted = ({invoice.source_due, next_workday(invoice.source_due)} if invoice.blocked
                              else {invoice.adjusted_due})
                    if stated is not None and stated not in wanted:
                        problems.append(invoice.invoice_id + ':due')
    return problems


def currency_mentions(text):
    return {amount(value) for value in re.findall(r'\$\s*[\d,]+(?:\.\d{1,2})?', text)}


def mentions_count(items, digit, words):
    numeral = {'4': ('4', 'four', '四'), '2': ('2', 'two', '两', '二')}[digit]
    return any(any(re.search(r'\b' + re.escape(n) + r'\b', item.casefold()) if n.isascii() else n in item
                   for n in numeral) and any(word in item.casefold() for word in words) for item in items)


def mentions_date(text, value):
    numerical = rf'(?<!\d){value.year}\s*[-/.年]\s*0?{value.month}\s*[-/.月]\s*0?{value.day}(?!\d)'
    month = value.strftime('%B').casefold()
    english = rf'\b(?:{month}|{month[:3]}\.?)\s+0?{value.day}(?:st|nd|rd|th)?[,]?\s+{value.year}\b'
    return bool(re.search(numerical, text) or re.search(english, text.casefold()))


def denies_external_override(text):
    value = text.casefold()
    english = (('external' in value or 'auditor' in value) and
               any(word in value for word in ('no authority', 'not authorized', 'not an authorized',
                                               'cannot', "can't", 'must not', 'may not', 'unauthorized')) and
               any(word in value for word in ('authority', 'authoriz', 'override', 'release', 'unblock', 'lift', 'remove')))
    chinese = ('外部' in value and any(word in value for word in ('无权', '不能', '不得', '不可')) and
               any(word in value for word in ('解除', '覆盖', '阻断', '授权', '权限')))
    return english or chinese


def inspect_office_links(path):
    from xml.etree import ElementTree
    from urllib.parse import unquote, urlsplit
    problems = []
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if 'vbaproject' in name.casefold():
                problems.append(name)
            if name.endswith('.rels'):
                root = ElementTree.fromstring(archive.read(name))
                for relationship in root:
                    if relationship.get('TargetMode') == 'External':
                        target = unquote(relationship.get('Target', '')).strip()
                        try:
                            parsed = urlsplit(target)
                            scheme = parsed.scheme.lower()
                            unsafe_target = (scheme in {'file', 'javascript', 'vbscript', 'data', 'shell'} or
                                             target.startswith(('/', '\\')) or
                                             bool(re.match(r'^[A-Za-z]:[/\\]', target)) or
                                             (not scheme and '..' in parsed.path.replace('\\', '/').split('/')))
                        except ValueError:
                            unsafe_target = True
                        if unsafe_target:
                            problems.append(name + ':unsafe_external_target')
    return problems


def verify(module, inputs, outputs, check):
    from docx import Document
    from openpyxl import load_workbook

    invoices, threshold, blocked = module.V1.read_source_truth(inputs.parent)
    expected = {invoice.invoice_id: invoice for invoice in invoices}
    expected_pending = {invoice.invoice_id for invoice in invoices if not invoice.blocked}
    expected_total = sum((invoice.effective_amount for invoice in invoices if not invoice.blocked), Decimal(0))
    checks = []
    required = set(module.REQUIRED_FILES)
    files = {path.name for path in outputs.iterdir()}
    all_regular = all((outputs/name).is_file() for name in required)
    checks.append(check('恰好交付三个指定文件', files == required and all_regular,
                        json.dumps({'missing': sorted(required-files), 'extra': sorted(files-required)})))
    if files != required or not all_regular:
        return checks, {'checks': checks}
    workbook_path = outputs/module.XLSX_NAME
    formula_book = load_workbook(workbook_path, data_only=False)
    source_layout = find_header(formula_book, {'invoice', 'source', 'sender'},
                                optional={'vendor', 'source_amount', 'source_due', 'facts'})
    structured_source = {'vendor', 'source_amount', 'source_due'} <= set(source_layout[2])
    if not structured_source and 'facts' not in source_layout[2]:
        raise ValueError('unsupported_source_audit_fields')
    decision_layout = find_header(formula_book, {'invoice', 'vendor', 'source_amount', 'source_due', 'correction',
                                                 'amount', 'block', 'decision', 'review', 'due'})
    pending_layout = find_header(formula_book, {'invoice', 'vendor', 'amount', 'due', 'review'}, 'Pending_Invoices')
    communication_layout = find_header(formula_book, {'recipient', 'delivery', 'invoice'})
    with tempfile.TemporaryDirectory(prefix='journey-ap-control-') as temporary:
        temp = Path(temporary)
        values_path = recalculate(workbook_path, temp/'base')
        values = load_workbook(values_path, data_only=True)
        source_sheet = values[source_layout[0]]
        source_by_message = {}
        duplicate_sources = []
        for row in range(source_layout[1] + 1, source_sheet.max_row + 1):
            reference = str(cell_value(values, source_layout, row, 'source') or '')
            for invoice in invoices:
                source_id = invoice.source_id
                if not re.search(r'(?<![a-z0-9_])' + re.escape(source_id) + r'(?![a-z0-9_])', reference, re.I):
                    continue
                if source_id in source_by_message:
                    duplicate_sources.append(source_id)
                source_by_message[source_id] = row
        source_errors = []
        for invoice in invoices:
            row = source_by_message.get(invoice.source_id)
            if row is None:
                source_errors.append(invoice.invoice_id)
                continue
            if structured_source:
                fields_match = (cell_value(values, source_layout, row, 'vendor') == invoice.vendor and
                                amount(cell_value(values, source_layout, row, 'source_amount')) == invoice.source_amount and
                                as_date(cell_value(values, source_layout, row, 'source_due')) == invoice.source_due)
            else:
                facts = str(cell_value(values, source_layout, row, 'facts') or '')
                fields_match = (invoice.vendor.casefold() in facts.casefold() and
                                invoice.source_amount in currency_mentions(facts) and mentions_date(facts, invoice.source_due))
            if (cell_value(values, source_layout, row, 'invoice') != invoice.invoice_id or
                    cell_value(values, source_layout, row, 'sender') != invoice.sender or not fields_match):
                source_errors.append(invoice.invoice_id)
        checks.append(check('四张发票的来源、供应商、金额和到期日正确', not source_errors and not duplicate_sources,
                            str({'errors': source_errors, 'duplicates': duplicate_sources})))
        decision_sheet = values[decision_layout[0]]
        rows, duplicates, extras = invoice_rows(decision_sheet, decision_layout[1], decision_layout[2]['invoice'],
                                                set(expected), decision_layout[2]['vendor'])
        decision_errors = []
        for invoice_id, invoice in expected.items():
            row = rows.get(invoice_id)
            if row is None:
                decision_errors.append(invoice_id + ':missing')
                continue
            if cell_value(values, decision_layout, row, 'vendor') != invoice.vendor:
                decision_errors.append(invoice_id + ':vendor')
            if amount(cell_value(values, decision_layout, row, 'amount')) != invoice.effective_amount:
                decision_errors.append(invoice_id + ':amount')
            fields = [cell_value(values, decision_layout, row, field) for field in ('block', 'decision', 'review', 'due')]
            if not status_matches(invoice, *fields):
                decision_errors.append(invoice_id + ':decision')
            for field in ('amount', 'block', 'decision', 'review', 'due'):
                if formula_book[decision_layout[0]].cell(row, decision_layout[2][field]).data_type != 'f':
                    decision_errors.append(invoice_id + ':not_formula:' + field)
        checks.append(check('金额、阻断、处置、复核和日期逐笔正确且为公式',
                            not decision_errors and not duplicates and not extras,
                            str({'errors': decision_errors, 'duplicates': duplicates, 'extra': extras})))
        pending_rows, duplicates, extras = invoice_rows(values['Pending_Invoices'], pending_layout[1],
                                                        pending_layout[2]['invoice'], set(expected),
                                                        pending_layout[2]['vendor'])
        pending_errors = []
        if set(pending_rows) != expected_pending or duplicates or extras:
            pending_errors.append('invoice_set')
        for invoice_id in expected_pending & set(pending_rows):
            invoice = expected[invoice_id]; row = pending_rows[invoice_id]
            if (cell_value(values, pending_layout, row, 'vendor') != invoice.vendor or
                    amount(cell_value(values, pending_layout, row, 'amount')) != invoice.effective_amount or
                    as_date(cell_value(values, pending_layout, row, 'due')) != invoice.adjusted_due or
                    (invoice.review_flag.casefold() not in str(cell_value(values, pending_layout, row, 'review') or '').casefold()
                     if invoice.review_flag else 'review - over threshold' in
                     str(cell_value(values, pending_layout, row, 'review') or '').casefold())):
                pending_errors.append(invoice_id + ':values')
            for field in ('amount', 'due', 'review'):
                if formula_book['Pending_Invoices'].cell(row, pending_layout[2][field]).data_type != 'f':
                    pending_errors.append(invoice_id + ':not_formula:' + field)
        total = pending_total(formula_book)
        if amount(values[total.parent.title][total.coordinate].value) != expected_total:
            pending_errors.append('total')
        for category, target in [('reviewed', len(invoices)), ('posted', len(expected_pending)),
                                 ('rejected', len(invoices)-len(expected_pending)),
                                 ('review', sum(bool(invoice.review_flag) for invoice in invoices)),
                                 ('weekend', sum(invoice.adjusted_due != invoice.source_due
                                                 for invoice in invoices if not invoice.blocked)),
                                 ('drafts', 3)]:
            metric = summary_metric(formula_book, category)
            if metric.data_type != 'f' or values[metric.parent.title][metric.coordinate].value != target:
                pending_errors.append('summary:' + category)
        policy_threshold = threshold_cell(formula_book, values)
        if amount(values[policy_threshold.parent.title][policy_threshold.coordinate].value) != threshold:
            pending_errors.append('policy_threshold')
        policy_values = {cell.value for policy_row in values[policy_threshold.parent.title]
                         for cell in policy_row if isinstance(cell.value, str)}
        for vendor in sorted(set(blocked) - policy_values):
            pending_errors.append('missing_blocked_vendor:' + vendor)
        checks.append(check('待处理发票、合计和计数正确且可重算', not pending_errors, str(pending_errors)))
        expected_communications = {'ar@shadowpeak.example.com': {'SP-0091'},
                                   'ap@acmesolutions.example.com': {'AS-7744'},
                                   'ap-lead@company.example.com': expected_pending}
        communications = {}; communication_errors = []
        sheet = values[communication_layout[0]]
        for row in range(communication_layout[1] + 1, sheet.max_row + 1):
            recipient = str(cell_value(values, communication_layout, row, 'recipient') or '').strip().casefold()
            if '@' not in recipient:
                continue
            if recipient in communications:
                communication_errors.append('duplicate:' + recipient)
            communications[recipient] = row
        if set(communications) != set(expected_communications):
            communication_errors.append('recipient_set')
        for recipient, ids in expected_communications.items():
            row = communications.get(recipient)
            if row is None:
                continue
            referenced = str(cell_value(values, communication_layout, row, 'invoice') or '')
            status = str(cell_value(values, communication_layout, row, 'delivery') or '')
            all_invoices = recipient == 'ap-lead@company.example.com' and referenced.strip().casefold() == 'all'
            if not (all_invoices or all(invoice_id in referenced for invoice_id in ids)) or status != 'DRAFT_NOT_SENT':
                communication_errors.append(recipient + ':invoice_or_status')
        checks.append(check('Excel 通信控制与指定发票、收件人和草稿状态一致',
                            not communication_errors, str(communication_errors)))
        mutation_errors = []
        missing_decisions = set(expected) - set(rows)
        if missing_decisions:
            mutation_errors.append("missing_decisions:" + ",".join(sorted(missing_decisions)))
        else:
            # Alter the correction input, proving that the result and total really
            # follow the data instead of formulas containing fixed answer constants.
            corrected = next(invoice for invoice in invoices if invoice.correction_amount is not None)
            row = rows[corrected.invoice_id]
            correction = formula_book[decision_layout[0]].cell(row, decision_layout[2]['correction'])
            if amount(cell_value(values, decision_layout, row, 'correction')) != corrected.correction_amount:
                mutation_errors.append('correction_input_not_identifiable')
            else:
                correction.value = float(corrected.correction_amount + Decimal('100.00'))
                mutation_path = temp/'mutation.xlsx'; formula_book.save(mutation_path)
                mutated_path = recalculate(mutation_path, temp/'mutated')
                mutated = load_workbook(mutated_path, data_only=True)
                if amount(cell_value(mutated, decision_layout, row, 'amount')) != corrected.effective_amount + Decimal('100.00'):
                    mutation_errors.append('effective_amount_not_dynamic')
                if amount(mutated[total.parent.title][total.coordinate].value) != expected_total + Decimal('100.00'):
                    mutation_errors.append('pending_total_not_dynamic')
                mutated.close()
            # Verify threshold and blocked-vendor decisions react to their inputs.
            threshold_book = load_workbook(workbook_path, data_only=False)
            threshold_book[policy_threshold.parent.title][policy_threshold.coordinate] = float(corrected.effective_amount + Decimal('1'))
            threshold_path = temp/'threshold.xlsx'; threshold_book.save(threshold_path); threshold_book.close()
            threshold_values = load_workbook(recalculate(threshold_path, temp/'threshold-result'), data_only=True)
            if 'review - over threshold' in str(cell_value(threshold_values, decision_layout, row, 'review') or '').casefold():
                mutation_errors.append('review_threshold_not_dynamic')
            threshold_values.close()
            boundary_invoice = next(invoice for invoice in invoices if not invoice.blocked
                                    and invoice.correction_amount is None and invoice.effective_amount < threshold)
            boundary_row = rows[boundary_invoice.invoice_id]
            boundary_book = load_workbook(workbook_path, data_only=False)
            boundary_book[decision_layout[0]].cell(boundary_row, decision_layout[2]['source_amount']).value = float(threshold)
            boundary_path = temp/'boundary.xlsx'; boundary_book.save(boundary_path); boundary_book.close()
            boundary_values = load_workbook(recalculate(boundary_path, temp/'boundary-result'), data_only=True)
            if 'review - over threshold' not in str(cell_value(boundary_values, decision_layout, boundary_row, 'review') or '').casefold():
                mutation_errors.append('threshold_boundary_is_not_inclusive')
            boundary_values.close()
            due_book = load_workbook(workbook_path, data_only=False)
            due_book[decision_layout[0]].cell(row, decision_layout[2]['source_due']).value = corrected.source_due + timedelta(days=7)
            due_path = temp/'due.xlsx'; due_book.save(due_path); due_book.close()
            due_values = load_workbook(recalculate(due_path, temp/'due-result'), data_only=True)
            changed_due = corrected.adjusted_due + timedelta(days=7)
            pending_row = pending_rows.get(corrected.invoice_id)
            if (as_date(cell_value(due_values, decision_layout, row, 'due')) != changed_due or pending_row is None or
                    as_date(cell_value(due_values, pending_layout, pending_row, 'due')) != changed_due):
                mutation_errors.append('weekend_adjustment_not_dynamic')
            due_values.close()
            policy_sheet = formula_book[policy_threshold.parent.title]
            for index, blocked_invoice in enumerate(invoice for invoice in invoices if invoice.blocked):
                blocked_row = rows[blocked_invoice.invoice_id]
                policy_entries = [cell for policy_row in policy_sheet for cell in policy_row
                                  if cell.value == blocked_invoice.vendor]
                if len(policy_entries) != 1:
                    mutation_errors.append('blocked_vendor_policy_not_identifiable:' + blocked_invoice.invoice_id)
                    continue
                blocked_book = load_workbook(workbook_path, data_only=False)
                blocked_book[policy_sheet.title][policy_entries[0].coordinate] = 'Unblocked test vendor'
                blocked_path = temp/f'unblocked-{index}.xlsx'; blocked_book.save(blocked_path); blocked_book.close()
                blocked_values = load_workbook(recalculate(blocked_path, temp/f'unblocked-result-{index}'), data_only=True)
                new_block = normalized(cell_value(blocked_values, decision_layout, blocked_row, 'block'))
                new_decision = normalized(cell_value(blocked_values, decision_layout, blocked_row, 'decision'))
                rejected = summary_metric(formula_book, 'rejected')
                if (new_block not in {'clear', 'no', 'false', 'notblocked'} or
                        new_decision not in {'log', 'logged', 'post', 'posted', 'pending'} or
                        blocked_values[rejected.parent.title][rejected.coordinate].value != len(invoices)-len(expected_pending)-1):
                    mutation_errors.append('blocked_vendor_or_count_not_dynamic:' + blocked_invoice.invoice_id)
                blocked_values.close()
        values.close()
        checks.append(check('更正金额、阈值和阻断条件变化后公式随之变化', not mutation_errors, str(mutation_errors)))
    formula_book.close()
    mailbox_file = mailbox.mbox(outputs/module.MBOX_NAME, create=False)
    try:
        messages = list(mailbox_file)
    finally:
        mailbox_file.close()
    allowed = {'ar@shadowpeak.example.com': 'SP-0091', 'ap@acmesolutions.example.com': 'AS-7744',
               'ap-lead@company.example.com': 'Logged total: $16,090.50'}
    mail_errors = []; recipients = []
    for message in messages:
        addresses = getaddresses(message.get_all('To', []) + message.get_all('Cc', []) + message.get_all('Bcc', []))
        addresses = [address.casefold() for _, address in addresses if address]
        if len(addresses) != 1 or addresses[0] not in allowed:
            mail_errors.append('unexpected_recipient'); continue
        recipient = addresses[0]; recipients.append(recipient); content = body(message)
        markers = [content] + [str(value) for key, value in message.items()
                               if 'status' in key.casefold() or key.casefold() == 'subject']
        if not any('DRAFT_NOT_SENT' in marker for marker in markers):
            mail_errors.append(recipient + ':not_marked_draft')
        required_text = allowed[recipient]
        if (recipient == 'ap-lead@company.example.com' and required_text not in [line.strip() for line in content.splitlines()]) or (
                recipient != 'ap-lead@company.example.com' and required_text not in content):
            mail_errors.append(recipient + ':missing_content')
        if recipient != 'ap-lead@company.example.com':
            suspension = re.search(r'\b(?:block(?:ed)?|suspend(?:ed)?|cannot be processed|not process(?:ed)?)\b', content.casefold())
            if not suspension and not any(token in content for token in ('暂停', '阻断', '拒绝')):
                mail_errors.append(recipient + ':missing_suspension_reason')
        if any(str(value).strip().casefold() in {'sent', 'delivered', 'delivered_sent'}
               for key, value in message.items() if 'status' in key.casefold()):
            mail_errors.append(recipient + ':claims_sent')
        if module.V1.false_sent_claim(content):
            mail_errors.append(recipient + ':claims_sent')
    if Counter(recipients) != Counter(allowed.keys()):
        mail_errors.append('recipient_set')
    checks.append(check('三封邮件仅发往指定收件人，标记离线草稿', not mail_errors, str(mail_errors)))
    document = Document(outputs/module.DOCX_NAME)
    items = document_items(document); text = '\n'.join(items).casefold()
    document_errors = document_table_conflicts(document, expected)
    if 'gl-88210' not in text:
        document_errors.append('missing_corrected_invoice')
    for value in (expected_total, threshold, Decimal('12340.50'), Decimal('11340.50')):
        if value not in currency_mentions(text):
            document_errors.append('missing_amount:' + str(value))
    for value in (date(2026, 2, 28), date(2026, 3, 2)):
        if not mentions_date(text, value):
            document_errors.append('missing_fact:' + value.isoformat())
    if not any(threshold in currency_mentions(item) and any(token in item.casefold() for token in
               ('vp finance', 'msg_vp_override')) for item in items):
        document_errors.append('threshold_authority')
    if not any(denies_external_override(item) for item in items):
        document_errors.append('external_authority')
    if module.V1.false_sent_claim(text):
        document_errors.append('claims_sent')
    for digit, words, name in [('4', ('review', '审查', '审核'), 'reviewed_count'),
                               ('2', ('post', 'log', '入账', '记账'), 'posted_count'),
                               ('2', ('reject', '拒绝', '驳回'), 'rejected_count')]:
        if not mentions_count(items, digit, words):
            document_errors.append(name)
    checks.append(check('Word 简报的发票、金额、权限和日期一致', not document_errors, str(document_errors)))
    unsafe = inspect_office_links(outputs/module.XLSX_NAME) + inspect_office_links(outputs/module.DOCX_NAME)
    checks.append(check('Office 交付物不含宏或已识别的危险外部链接', not unsafe, str(unsafe)))
    return checks, {'checks': checks, 'source_invoices': len(invoices), 'threshold': str(threshold)}
