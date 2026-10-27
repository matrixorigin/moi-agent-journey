"""Check summary values against labels and data in the original child sheets."""
import math
import subprocess
import tempfile
from pathlib import Path


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def source_cells(sheet):
    volume = []
    total = []
    for row in sheet:
        for cell in row:
            if not isinstance(cell.value, str):
                continue
            label = cell.value.strip()
            if label.startswith('砼体积'):
                volume.append(cell)
            if label == '总重量合计':
                total.append(cell)
    if len(volume) != 1 or len(total) != 1:
        raise ValueError(f'{sheet.title}:ambiguous_source_labels:{len(volume)}/{len(total)}')
    volume_cell = sheet.cell(volume[0].row + 1, volume[0].column)
    # A total may be displayed next to its label, or under a column header.
    candidates = [cell for cell in sheet[total[0].row] if numeric(cell.value)]
    if len(candidates) == 1:
        total_cell = candidates[0]
    elif not candidates:
        total_cell = sheet.cell(total[0].row + 1, total[0].column)
    else:
        raise ValueError(sheet.title + ':ambiguous_total_values')
    if not numeric(volume_cell.value) or not numeric(total_cell.value):
        raise ValueError(sheet.title + ':source_value_not_numeric')
    return volume_cell, total_cell


def value_equal(actual, expected):
    return numeric(actual) and numeric(expected) and math.isclose(actual, expected, rel_tol=1e-9, abs_tol=1e-8)


def verify(inputs, outputs, check):
    from openpyxl import load_workbook

    source_path = inputs / '子表数据提取到汇总表.xlsx'
    output_path = outputs / 'summary.xlsx'
    source = load_workbook(source_path, data_only=False)
    source_values = load_workbook(source_path, data_only=True)
    submitted = load_workbook(output_path, data_only=False)
    checks = [check('保留工作表及顺序', source.sheetnames == submitted.sheetnames)]
    if source.sheetnames != submitted.sheetnames:
        return checks, {'checks': checks}
    source_summary = source['汇总']
    mappings = [(row, source_summary.cell(row, 3).value)
                for row in range(3, source_summary.max_row + 1)
                if source_summary.cell(row, 3).value is not None]
    children = [name for name in source.sheetnames if name != '汇总']
    names = [name for _, name in mappings]
    valid_mapping = len(names) == len(set(names)) and set(names) == set(children)
    if not valid_mapping:
        raise ValueError('source_summary_does_not_map_every_child_once')
    with tempfile.TemporaryDirectory(prefix='journey-summary-') as tmp:
        directory = Path(tmp)
        calculated_dir = directory / 'calculated'
        calculated_dir.mkdir()
        result = subprocess.run(['libreoffice', '-env:UserInstallation=' + (directory/'profile').as_uri(),
                                 '--headless', '--convert-to', 'xlsx', '--outdir', str(calculated_dir),
                                 str(output_path)], capture_output=True, text=True, timeout=120)
        calculated_path = calculated_dir / output_path.name
        if result.returncode or not calculated_path.exists():
            checks.append(check('工作簿可重新计算', False, (result.stdout + result.stderr)[-2000:]))
            return checks, {'checks': checks}
        calculated = load_workbook(calculated_path, data_only=True)
        errors = []
        for row, name in mappings:
            volume, total = source_cells(source_values[name])
            for col, expected in ((6, volume), (7, total)):
                actual = calculated['汇总'].cell(row, col).value
                if not value_equal(actual, expected.value):
                    errors.append({'summary': calculated['汇总'].cell(row, col).coordinate,
                                   'source': name + '!' + expected.coordinate,
                                   'expected': expected.value, 'actual': actual})
        checks.append(check('每个子表的体积和总重量正确汇总', not errors,
                            str({'children': len(children), 'mismatches': errors[:12]})))
        calculated.close()
    targets = {(row, column) for row, _ in mappings for column in (6, 7)}
    differences = []
    for sheet in source:
        actual = submitted[sheet.title]
        if set(map(str, sheet.merged_cells.ranges)) != set(map(str, actual.merged_cells.ranges)):
            differences.append(sheet.title + ':merged_cells_changed')
        max_row, max_col = max(sheet.max_row, actual.max_row), max(sheet.max_column, actual.max_column)
        for row in range(1, max_row + 1):
            for col in range(1, max_col + 1):
                if sheet.title == '汇总' and (row, col) in targets:
                    continue
                if sheet.cell(row, col).value != actual.cell(row, col).value:
                    differences.append(sheet.title + '!' + sheet.cell(row, col).coordinate)
    checks.append(check('保留非目标内容、原有公式和合并单元格', not differences,
                        str({'changed': differences[:12], 'count': len(differences)})))
    source.close(); source_values.close(); submitted.close()
    return checks, {'checks': checks, 'child_count': len(children)}
