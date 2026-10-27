"""Trusted verification, executed with no network or QA credentials in Docker."""
import argparse
import copy
import importlib.util
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from journey.evidence import sha256


def module_at(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def check(rule, passed, detail=""):
    return {"rule": rule, "passed": bool(passed), "detail": detail}


class PytestCapture:
    """Keep inapplicable checks distinct from success and failure."""
    def __init__(self):
        self.allowed_skips = set()
        self.checks = []
        self.skipped = []
        self.unexpected_skips = []

    def record(self, report):
        if report.skipped:
            item = {"rule": report.nodeid, "detail": str(report.longrepr)}
            (self.skipped if report.nodeid in self.allowed_skips else self.unexpected_skips).append(item)
        elif report.when == "call" or report.failed:
            self.checks.append(check(report.nodeid, report.passed,
                                     str(report.longrepr) if report.failed else report.outcome))

    def result(self, code):
        if code not in (0, 1):
            raise RuntimeError("oracle_collection_or_runtime_error:" + str(code))
        if self.unexpected_skips:
            raise RuntimeError("oracle_unexpected_skipped_checks:" +
                               ",".join(item["rule"] for item in self.unexpected_skips))
        if not self.checks:
            raise RuntimeError("oracle_executed_no_checks")
        if code == 1 and not any(not item["passed"] for item in self.checks):
            raise RuntimeError("oracle_failed_without_recorded_check")
        return self.checks, {"exit_code": int(code), "checks": self.checks, "skipped": self.skipped}


def adapt_skillsbench_states(case, module, inputs):
    if case["id"] == "J1-08":
        from pypdf import PdfReader
        pages = PdfReader(str(inputs / "population.pdf")).pages
        if not any("Other Territories" in (page.extract_text() or "") for page in pages):
            raise RuntimeError("population_source_missing_other_territories")
        # The public population source explicitly includes this ninth region.
        module.VALID_STATES.add("Other Territories")


def render_outputs(outputs, destination):
    results = []
    destination.mkdir(parents=True, exist_ok=True)
    for path in sorted(outputs.rglob("*")):
        if path.suffix.lower() not in {".docx", ".pptx", ".xlsx"}:
            continue
        target = destination / path.relative_to(outputs).as_posix().replace("/", "_")
        target.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory() as profile:
            command = ["libreoffice", "-env:UserInstallation=" + Path(profile).as_uri(),
                       "--headless", "--convert-to", "pdf", "--outdir", str(target), str(path)]
            proc = subprocess.run(command, capture_output=True, text=True, timeout=90)
        pdf = target / (path.stem + ".pdf")
        good = proc.returncode == 0 and pdf.is_file() and pdf.stat().st_size > 0
        results.append(check("可渲染：" + path.relative_to(outputs).as_posix(), good,
                             (proc.stdout + proc.stderr)[-3000:]))
    return results


def officeval(case, oracle, inputs, outputs):
    result = module_at(oracle / "judge.py", "officeval_oracle").evaluate(str(outputs))
    if case["id"] in {"J1-03", "J1-06", "J1-12"}:
        from . import statement, summary_sheet, table_chart
        verifier = {"J1-03": statement.verify, "J1-06": summary_sheet.verify,
                    "J1-12": table_chart.verify}[case["id"]]
        checks, raw = verifier(inputs, outputs, check)
        raw["reference_upstream_result"] = result
        return checks, raw
    checks = [check("上游文件与结构", result.get("dim1_pass"), result.get("dim1_reason", ""))]
    for item in result["dim2_items"]:
        positive = item["max_delta"] > 0
        checks.append(check(item["rule"], item["hit"] if positive else not item["hit"], item.get("detail", "")))
    if result.get("status") != "ok":
        checks.append(check("文件可被检查器解析", False, str(result.get("error"))))
    return checks, result


def skills(case, oracle, inputs, outputs):
    if case["id"] == "J1-16":
        module = module_at(oracle / "judge.py", "organize_oracle")
        expected = {folder + "/" + name: name for name, folder in module.FILE_TO_SUBJECT.items()}
        actual = {p.relative_to(outputs).as_posix(): p for p in outputs.rglob("*") if p.is_file()}
        input_hashes = {a["path"]: a["sha256"] for a in case["assets"]}
        checks = [check("交付主题、文件名、完整集合", set(actual) == set(expected),
            json.dumps({"missing": sorted(set(expected)-set(actual)), "unexpected": sorted(set(actual)-set(expected))}))]
        for relative, name in expected.items():
            checks.append(check("字节保持：" + relative,
                                relative in actual and sha256(actual[relative]) == input_hashes[name]))
        return checks, {"checks": checks}
    import pytest
    with tempfile.TemporaryDirectory() as temporary:
        test = Path(temporary) / "test_upstream.py"
        shutil.copyfile(oracle / "judge.py", test)
        capture = PytestCapture()

        class Adapter:
            def pytest_collection_modifyitems(self, items):
                modules = {item.module for item in items}
                def remap(value):
                    if isinstance(value, (str, Path)) and str(value).startswith("/root/"):
                        relative = str(value)[6:]
                        destination = inputs if (inputs / relative).is_file() else outputs
                        path = destination / relative
                        return str(path) if isinstance(value, str) else path
                    if isinstance(value, list):
                        return [remap(v) for v in value]
                    return value
                for module in modules:
                    for name, value in list(vars(module).items()):
                        if name.isupper():
                            setattr(module, name, remap(value))
                    if case["id"] == "J1-04":
                        module.EXPECTED_FILE = oracle / "expected_output.json"
                    adapt_skillsbench_states(case, module, inputs)
                if case["id"] == "J1-08":
                    for item in items:
                        if (getattr(item, "originalname", None) == "test_pivot_col_field"
                                and item.callspec.params.get("col_field") is None):
                            capture.allowed_skips.add(item.nodeid)

            def pytest_runtest_logreport(self, report):
                capture.record(report)

        code = pytest.main(["-q", "-p", "no:cacheprovider", str(test)], plugins=[Adapter()])
        return capture.result(code)


def lbo(oracle, inputs, outputs):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    source = load_workbook(inputs / "09_01_input.xlsx", data_only=False)
    fixed = load_workbook(outputs / "09_01_repaired.xlsx", data_only=False)
    # The financial relation is the assertion; equivalent formulas are checked after recalculation.
    checks = [check("保留工作表和顺序", fixed.sheetnames == source.sheetnames)]
    if fixed.sheetnames != source.sheetnames:
        return checks, {"checks": checks}
    targets = {}
    for col in range(7, 15):
        c = get_column_letter(col)
        targets.update({("LBO", f"{c}66"): [f"{c}64", f"{c}75"],
                        ("LBO", f"{c}69"): [f"{c}66", f"{c}67", f"{c}68"],
                        ("LBO", f"{c}175"): [f"{c}136", f"{c}137", f"{c}145", f"{c}152"]})
    for col in range(10, 15):
        c = get_column_letter(col)
        targets[("Financials", f"{c}18")] = [f"{c}12", f"{c}15", f"{c}16", f"{c}17"]
    with tempfile.TemporaryDirectory() as directory:
        # LBO interest/debt is circular by design. Set a deterministic calculation
        # environment on a verification copy; never alter the submitted workbook.
        scenario = Path(directory)/'scenario'
        calculated_dir = Path(directory)/'calculated'
        scenario.mkdir(); calculated_dir.mkdir()
        calculation_copy = scenario/'09_01_repaired.xlsx'
        fixed.calculation.iterate=True
        fixed.calculation.iterateCount=1000
        fixed.calculation.iterateDelta=1e-8
        fixed.save(calculation_copy)
        result = subprocess.run(["libreoffice", "-env:UserInstallation=" + (Path(directory)/"profile").as_uri(),
            "--headless", "--convert-to", "xlsx", "--outdir", str(calculated_dir), str(calculation_copy)], capture_output=True, text=True, timeout=120)
        calculated = calculated_dir / "09_01_repaired.xlsx"
        if result.returncode or not calculated.exists():
            raise RuntimeError("lbo_recalculation_failed:" + result.stderr)
        values = load_workbook(calculated, data_only=True)
        for (sheet, cell), refs in targets.items():
            terms = [values[sheet][ref].value for ref in refs]
            actual = values[sheet][cell].value
            if not all(isinstance(term, (int, float)) for term in terms):
                checks.append(check(sheet+"!"+cell+" 依赖可重算",False,str(terms)))
                continue
            expected = terms[0] - sum(terms[1:]) if sheet == "Financials" else sum(terms)
            formula = fixed[sheet][cell].data_type == "f"
            checks.append(check(sheet+"!"+cell+" 的业务关系和原生公式", formula and isinstance(actual,(int,float)) and abs(actual-expected)<1e-7,
                                json.dumps({"actual":actual,"expected":expected})))
        for sheet in source:
            for row in sheet:
                for cell in row:
                    if (sheet.title, cell.coordinate) in targets:
                        continue
                    actual = fixed[sheet.title][cell.coordinate]
                    same_style = all(copy.copy(getattr(cell, key)) == copy.copy(getattr(actual, key))
                        for key in ['font','fill','border','alignment','protection','number_format'])
                    if cell.value != actual.value or not same_style:
                        checks.append(check("保留非目标单元格："+sheet.title+"!"+cell.coordinate,False))
            for row in values[sheet.title]:
                for cell in row:
                    if cell.data_type == "e":
                        checks.append(check("重算无错误："+sheet.title+"!"+cell.coordinate,False,str(cell.value)))
        values.close()
        # Change independent drivers to distinguish dynamic relations from formulas
        # that simply return the correct constant for this one fixed snapshot.
        drivers = {(sheet,ref) for (sheet,_),refs in targets.items() for ref in refs} - set(targets)
        for index,(sheet,ref) in enumerate(sorted(drivers)):
            fixed[sheet][ref] = 101 + index * 3
        fixed.save(calculation_copy)
        calculated.unlink()
        mutated = subprocess.run(["libreoffice", "-env:UserInstallation=" + (Path(directory)/"mutation-profile").as_uri(),
            "--headless", "--convert-to", "xlsx", "--outdir", str(calculated_dir), str(calculation_copy)],capture_output=True,text=True,timeout=120)
        if mutated.returncode or not calculated.exists():
            raise RuntimeError("lbo_mutation_recalculation_failed:"+mutated.stderr)
        values=load_workbook(calculated,data_only=True)
        for (sheet,cell),refs in targets.items():
            terms=[values[sheet][ref].value for ref in refs]
            actual=values[sheet][cell].value
            numeric=all(isinstance(term,(int,float)) for term in terms)
            expected=(terms[0]-sum(terms[1:]) if sheet=='Financials' else sum(terms)) if numeric else None
            checks.append(check(sheet+'!'+cell+' 输入变化后仍正确计算',numeric and isinstance(actual,(int,float)) and abs(actual-expected)<1e-7,
                                json.dumps({'actual':actual,'expected':expected})))
        values.close()
    checks.append(check("打开自动重算", fixed.calculation.calcMode in (None,"auto") and fixed.calculation.fullCalcOnLoad))
    source.close(); fixed.close()
    return checks,{"checks":checks}


def require_evo_checks_executed(case, result):
    if case["id"] != "J1-14":
        return
    criteria = {item["criterion_id"]: item for item in result.get("criterion_results", [])}
    evidence = criteria.get("R012", {}).get("evidence", "")
    if "Excel 工作表应为" in evidence:
        raise RuntimeError("upstream_oracle_incomplete: workbook layout caused invoice checks to be skipped")


def evo(case, oracle, inputs, outputs):
    if case["id"] == "J1-07":
        return lbo(oracle, inputs, outputs)
    module = module_at(oracle / "judge.py", "evo_oracle")
    result = module.evaluate(outputs)
    require_evo_checks_executed(case, result)
    checks=[]
    adjusted = {}
    if case["id"] == "J1-09":
        from .inventory import verify_dates
        criteria = {item["criterion_id"]: item for item in result["criterion_results"]}
        if "R001" not in criteria:
            raise RuntimeError("upstream_inventory_criteria_missing")
        if not criteria["R001"]["raw"]["passed"]:
            return [check("R001", False, criteria["R001"].get("evidence", ""))], result
        if not {"R006", "R007", "R008", "R011"} <= set(criteria):
            raise RuntimeError("upstream_inventory_criteria_missing")
        dates_ok, errors = verify_dates(module, inputs, outputs)
        needs_ok = criteria["R006"]["raw"]["passed"]
        schedule_ok = criteria["R008"]["raw"]["passed"]
        adjusted = {
            "R007": check("R007", dates_ok, json.dumps(errors, ensure_ascii=False)),
            "R011": check("R011", needs_ok and dates_ok and schedule_ok,
                          json.dumps({"needs": needs_ok, "dates": dates_ok, "schedule": schedule_ok})),
        }
    for item in result["criterion_results"]:
        # Full business completion requires every positive unit and no penalty, not a soft score.
        raw=item.get("raw",{})
        cid=item["criterion_id"]
        if cid in adjusted:
            checks.append(adjusted[cid])
            continue
        # The older judges normalize penalty scores to 1 when no violation occurs;
        # the v2 atomic judges instead report penalty severity (0 means no violation).
        if raw.get("type") == "penalty":
            passed = not raw["passed"]
        else:
            penalty = cid.startswith("P") or (case["id"]=="J1-14" and cid in {"R015","R016"})
            passed = item["score"] == (0 if penalty else 1)
        checks.append(check(cid,passed,item.get("evidence","")))
    if adjusted:
        result["journey_adapted_checks"] = list(adjusted.values())
    return checks,result


def pi_structural(case, outputs, session):
    checks=[]
    if case["id"] == "J1-17":
        if session.endswith("019"):
            data=json.loads((outputs/"merged_procurement_plan.json").read_text())
            required={"primary_procurement_move","immediate_slots","next_wave_gel_basics","conditional_backup_lab_support",
                "deferred_route_tied_confirmation","confirmation_action","fallback_if_no_reply","draft_corrections"}
            checks.append(check("采购计划独立决策字段",required.issubset(data)))
            slots=data.get('immediate_slots')
            valid=isinstance(slots,list) and all(isinstance(s,dict) and isinstance(s.get('item'),str) for s in slots)
            expected={'Target-gene primers','Sealing film','RNase-free filter tips','CCK-8 reagent','96-well cell plates','Carbomer 940','Triethanolamine','Propylene Glycol'}
            items=[s['item'] for s in slots] if valid else []
            checks.append(check('立即采购为指定的八种物料且无重复',valid and set(items)==expected and len(items)==len(set(items)),
                                json.dumps({'expected':sorted(expected),'actual':items})))
        else:
            import csv
            file=outputs/("gel_procurement_draft.csv" if session.endswith("016") else "qpcr_procurement_draft.csv")
            rows=list(csv.DictReader(file.open()))
            checks.append(check("修订后的 CSV 有可解析业务行",bool(rows)))
    else:
        for file in outputs.glob("*.md"):
            text=file.read_text()
            checks.append(check("非空 Markdown："+file.name,bool(text.strip())))
        if session.endswith("017"):
            text=(outputs/"Campaign_Outline.md").read_text()
            for heading in ['Co-branded Product','Target Audience Persona','Slogan','Promotion Strategy']:
                checks.append(check("方案业务模块："+heading,heading in text))
    return checks,{"checks":checks,"semantic_judge_required":True}


def preflight(outputs):
    import zipfile
    import fitz
    checks=[]
    for file in outputs.rglob('*'):
        if not file.is_file():continue
        try:
            if file.suffix.lower() in {'.docx','.xlsx','.pptx'}:
                with zipfile.ZipFile(file) as archive:
                    if archive.testzip() is not None:raise ValueError('OOXML CRC failure')
                    if '[Content_Types].xml' not in archive.namelist():raise ValueError('missing OOXML content types')
            elif file.suffix.lower()=='.json':json.loads(file.read_text())
            elif file.suffix.lower()=='.pdf':
                with fitz.open(file) as pdf:
                    if pdf.page_count<1:raise ValueError('empty PDF')
            checks.append(check('文件可解析：'+file.relative_to(outputs).as_posix(),True))
        except (ValueError, zipfile.BadZipFile, RuntimeError) as error:
            checks.append(check('文件可解析：'+file.relative_to(outputs).as_posix(),False,str(error)))
    return checks


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--case",required=True)
    parser.add_argument("--session",required=True)
    args=parser.parse_args()
    case=json.loads(Path(args.case).read_text())
    oracle,inputs,outputs=Path('/work/oracle'),Path('/work/inputs'),Path('/outputs')
    result={"status":"error","checks":[],"passed":False}
    try:
        checks=preflight(outputs)
        if any(not c['passed'] for c in checks):
            result={'status':'completed','checks':checks,'passed':False}
            Path('/verification/verdict.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
            return 0
        family=case['source']['benchmark']
        if family=='OfficeVal':checks,raw=officeval(case,oracle,inputs,outputs)
        elif family=='SkillsBench':checks,raw=skills(case,oracle,inputs,outputs)
        elif family=='π-Bench':checks,raw=pi_structural(case,outputs,args.session)
        elif family in {'EvoCowork','SpreadsheetBench 2 经 EvoCowork 改编'}:checks,raw=evo(case,oracle,inputs,outputs)
        else:raise ValueError('unsupported_benchmark_validator:'+family)
        rendered = Path('/verification/rendered')
        rendering = render_outputs(outputs,rendered)
        checks += rendering
        if case['id'] == 'J1-03' and all(item['passed'] for item in rendering):
            from .statement import verify_rendered
            checks.append(verify_rendered(outputs,rendered,check))
        result={"status":"completed","checks":checks,"upstream_result":raw,
                "passed":bool(checks) and all(c['passed'] for c in checks)}
    except Exception as error:
        result['reason']=type(error).__name__+":"+str(error)
    Path('/verification/verdict.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    return 0 if result['status']=='completed' else 1


if __name__=='__main__':
    raise SystemExit(main())
