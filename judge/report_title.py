"""Validate report title and abstract against their actual source requirements."""
from pathlib import Path


def styles(style):
    visited = set()
    while style is not None and style.style_id not in visited:
        visited.add(style.style_id)
        yield style
        style = style.base_style


def font_setting(run, paragraph, name):
    from docx.oxml.ns import qn

    def value(font, element):
        if name == 'name':
            rpr = element.rPr if element is not None else None
            if rpr is not None and rpr.rFonts is not None:
                east_asia = rpr.rFonts.get(qn('w:eastAsia'))
                if east_asia:
                    return east_asia
        return getattr(font, name)

    actual = value(run.font, run._element)
    if actual is not None:
        return actual
    for style in list(styles(run.style)) + list(styles(paragraph.style)):
        actual = value(style.font, style._element)
        if actual is not None:
            return actual
    return None


def paragraph_setting(paragraph, name):
    value = getattr(paragraph.paragraph_format, name)
    if value is not None:
        return value
    for style in styles(paragraph.style):
        value = getattr(style.paragraph_format, name)
        if value is not None:
            return value
    return None


def run_font_matches(paragraph, allowed_names, expected_size, bold=None):
    runs = [run for run in paragraph.runs if run.text.strip()]
    if not runs:
        return False
    for run in runs:
        font = font_setting(run, paragraph, 'name')
        size = font_setting(run, paragraph, 'size')
        actual_bold = font_setting(run, paragraph, 'bold')
        if (not isinstance(font, str) or font.casefold() not in allowed_names or size is None or
                abs(size.pt - expected_size) > 0.1 or (bold is not None and actual_bold is not bold)):
            return False
    return True


def source_title(inputs):
    from docx import Document
    document = Document(str(inputs/'公共书房运营规划书.docx'))
    title = next((paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()), '')
    if not title:
        raise ValueError('source_report_title_missing')
    return title


def verify(inputs, outputs, check):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

    title = source_title(inputs)
    doc = Document(str(outputs/'report.docx'))
    titles = [(index, paragraph) for index, paragraph in enumerate(doc.paragraphs) if paragraph.text.strip() == title]
    abstracts = [(index, paragraph) for index, paragraph in enumerate(doc.paragraphs)
                 if ''.join(paragraph.text.split()) == '摘要']
    title_ok = len(titles) == 1
    if title_ok:
        paragraph = titles[0][1]
        title_ok = (run_font_matches(paragraph, {'宋体', 'simsun'}, 16, True) and
                    paragraph_setting(paragraph, 'alignment') == WD_ALIGN_PARAGRAPH.CENTER)
        rule = paragraph_setting(paragraph, 'line_spacing_rule')
        spacing = paragraph_setting(paragraph, 'line_spacing')
        title_ok = title_ok and (rule == WD_LINE_SPACING.SINGLE or
                                 isinstance(spacing, float) and abs(spacing - 1.0) < 0.001)
    abstract_ok = (len(abstracts) == 1 and len(titles) == 1 and titles[0][0] < abstracts[0][0] and
                   run_font_matches(abstracts[0][1], {'黑体', 'simhei'}, 14))
    return [check('源报告标题：宋体三号、加粗、居中及单倍行距', title_ok,
                  str({'source_title': title, 'matches': len(titles)})),
            check('摘要标题：黑体四号且位于标题之后', abstract_ok,
                  str({'matches': len(abstracts)}))]


def verify_rendered(inputs, rendered, check):
    from pypdf import PdfReader

    title = source_title(inputs)
    pdf = PdfReader(str(Path(rendered)/'report.docx'/'report.pdf'))
    first_page = ''.join((pdf.pages[0].extract_text() or '').split()) if pdf.pages else ''
    return check('源报告标题渲染在第一页', ''.join(title.split()) in first_page)
