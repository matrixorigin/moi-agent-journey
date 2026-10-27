"""Verify a bank statement against the actual PDF and existing Word template."""
from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation
import re


AMOUNT = r"[\d,]+\.\d{2}"
DATE_PREFIX = r"(\d{1,2})\s+([A-Za-z]{3})\."
TRANSACTION = re.compile(rf"^{DATE_PREFIX}\s+(.+?)\s+({AMOUNT})\s+({AMOUNT})$")
OPENING = re.compile(rf"^{DATE_PREFIX}\s+Balance forward\s+\$?({AMOUNT})$", re.I)
CLOSING = re.compile(rf"^{DATE_PREFIX}\s+Closing totals\s+({AMOUNT})\s+({AMOUNT})\s+({AMOUNT})$", re.I)
HEADER = ('date', 'description', 'withdrawals', 'deposits', 'balance')


def text(value):
    return ' '.join(str(value or '').split()).casefold()


def money(value):
    value = str(value or '').strip().replace(',', '').replace('$', '')
    if not value:
        return None
    try:
        parsed = Decimal(value)
        if not parsed.is_finite():
            raise ValueError('invalid_amount:' + value)
        return parsed
    except InvalidOperation as error:
        raise ValueError('invalid_amount:' + value) from error


def day_month(value):
    value = re.sub(r'\.', '', ' '.join(str(value).split()))
    for fmt in ('%d %b', '%b %d', '%d %b %Y', '%b %d %Y'):
        try:
            parsed = datetime.strptime(value, fmt)
            return parsed.month, parsed.day
        except ValueError:
            pass
    raise ValueError('invalid_date:' + value)


def pdf_records(content):
    """Derive transactions and reconcile them to the statement's own totals."""
    opening = closing = None
    balance = None
    records = []
    for line in content.splitlines():
        line = ' '.join(line.split())
        match = OPENING.fullmatch(line)
        if match:
            if opening is not None:
                raise ValueError('multiple_opening_balances')
            opening = balance = money(match[3])
            continue
        match = CLOSING.fullmatch(line)
        if match:
            if closing is not None:
                raise ValueError('multiple_closing_totals')
            closing = tuple(money(match[i]) for i in (3, 4, 5))
            continue
        match = TRANSACTION.fullmatch(line)
        if not match or balance is None:
            continue
        date = day_month(match[1] + ' ' + match[2])
        description = text(match[3])
        amount, next_balance = money(match[4]), money(match[5])
        delta = next_balance - balance
        if delta == -amount:
            withdrawal, deposit = amount, None
        elif delta == amount:
            withdrawal, deposit = None, amount
        else:
            raise ValueError('pdf_transaction_does_not_reconcile:' + description)
        records.append((date, description, withdrawal, deposit, next_balance))
        balance = next_balance
    withdrawals = sum((r[2] or Decimal(0) for r in records), Decimal(0))
    deposits = sum((r[3] or Decimal(0) for r in records), Decimal(0))
    if opening is None or closing is None or not records or closing != (withdrawals, deposits, balance):
        raise ValueError('pdf_totals_do_not_reconcile')
    return records, closing


def normalize_row(cells):
    if len(cells) != 5:
        raise ValueError('invalid_column_count')
    date, description, withdrawal, deposit, balance = cells
    withdrawn, deposited, balance = money(withdrawal), money(deposit), money(balance)
    if withdrawn is not None:
        withdrawn = abs(withdrawn)
    if deposited is not None and deposited < 0:
        raise ValueError('negative_deposit')
    if (withdrawn is None) == (deposited is None) or balance is None:
        raise ValueError('invalid_transaction_amounts')
    if not text(description):
        raise ValueError('missing_description')
    return day_month(date), text(description), withdrawn, deposited, balance


def record_difference(expected, actual):
    wanted, received = Counter(expected), Counter(actual)
    return wanted - received, received - wanted


def explicit_year_errors(document, source_records, source_year):
    source_keys = {(row[0], row[1]) for row in source_records}
    errors = []
    for table in document.tables:
        if not table.rows or tuple(text(c.text) for c in table.rows[0].cells) != HEADER:
            continue
        for row in table.rows[1:]:
            if len(row.cells) != 5:
                continue
            date, description = row.cells[0].text, text(row.cells[1].text)
            year = re.search(r'\b(\d{4})\b', date)
            if not year:
                continue
            try:
                key = day_month(date), description
            except ValueError:
                continue
            if key in source_keys and int(year.group(1)) != source_year:
                errors.append({'date': date, 'description': description, 'expected_year': source_year})
    return errors


def document_rows(document):
    rows, heights, bad, closing_rows = [], [], [], []
    found_table = False
    for table in document.tables:
        if not table.rows or tuple(text(c.text) for c in table.rows[0].cells) != HEADER:
            continue
        found_table = True
        for row in table.rows[1:]:
            values = [c.text for c in row.cells]
            if not any(text(v) for v in values):
                continue
            if len(values) == 5 and text(values[1]) == 'closing totals':
                closing_rows.append(values)
                heights.append((row.height, row.height_rule))
                continue
            try:
                rows.append(normalize_row(values))
                heights.append((row.height, row.height_rule))
            except ValueError as error:
                bad.append(str(error))
    return found_table, rows, heights, bad, closing_rows


def rendered_cells_visible(document, rendered_text):
    rendered_text = text(rendered_text)
    missing = []
    for table in document.tables:
        if not table.rows or tuple(text(c.text) for c in table.rows[0].cells) != HEADER:
            continue
        for row_number, row in enumerate(table.rows[1:], 1):
            for column, cell in enumerate(row.cells, 1):
                content = text(cell.text)
                if content and content not in rendered_text:
                    missing.append({'row': row_number, 'column': column, 'text': content})
    return missing


def verify_rendered(outputs, rendered, check):
    from docx import Document
    from pypdf import PdfReader

    doc = Document(str(outputs / 'statement.docx'))
    pdf = PdfReader(str(rendered / 'statement.docx' / 'statement.pdf'))
    missing = rendered_cells_visible(doc, '\n'.join(page.extract_text() or '' for page in pdf.pages))
    return check('记录单元格渲染后可见', not missing, str({'missing': missing[:8]}))


def verify(inputs, outputs, check):
    from docx import Document
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from pypdf import PdfReader

    pdf = PdfReader(str(inputs / 'Cedarline_Statement_2025DEC_reordered.pdf'))
    pdf_text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
    periods = set(re.findall(r'\b\d{1,2}\s+[A-Za-z]{3}\.\s+(\d{4})\s+to\s+'
                             r'\d{1,2}\s+[A-Za-z]{3}\.\s+(\d{4})\b', pdf_text, re.I))
    if len(periods) != 1 or next(iter(periods))[0] != next(iter(periods))[1]:
        raise ValueError('source_statement_period_missing_or_ambiguous')
    year = int(next(iter(periods))[0])
    source, closing = pdf_records(pdf_text)
    template = Document(str(inputs / 'public_ready_statement.docx'))
    template_found, original, _, template_bad, _ = document_rows(template)
    if not template_found or template_bad:
        raise ValueError('template_records_unreadable:' + ';'.join(template_bad))
    document = Document(str(outputs / 'statement.docx'))
    found, actual, heights, malformed, closing_rows = document_rows(document)
    year_errors = explicit_year_errors(document, source, year)
    missing, unexpected = record_difference(source + original, actual)
    def sample(counter):
        return [str(record) + ' x' + str(count) for record, count in list(counter.items())[:8]]
    totals_good = len(closing_rows) <= 1
    if closing_rows:
        values = closing_rows[0]
        try:
            totals_good = totals_good and tuple(abs(money(values[i])) for i in (2, 3, 4)) == closing
        except (TypeError, ValueError):
            totals_good = False
    checks = [
        check('原生可编辑记录表', found),
        check('PDF 与模板记录完整且无重复', not missing and not unexpected and not malformed and not year_errors,
              str({'source_transactions': len(source), 'template_records': len(original),
                   'missing': sample(missing), 'unexpected': sample(unexpected), 'malformed': malformed[:8],
                   'wrong_year': year_errors[:8]})),
        check('如提供结算合计则与 PDF 一致', totals_good),
        check('记录行使用统一的固定行高', bool(heights) and len(set(heights)) == 1
              and heights[0][0] is not None and heights[0][1] == WD_ROW_HEIGHT_RULE.EXACTLY,
              str({'record_rows': len(heights), 'distinct_heights_and_rules': len(set(heights))})),
    ]
    return checks, {'checks': checks, 'source_transactions': len(source), 'template_records': len(original)}
