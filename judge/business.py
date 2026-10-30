"""Business-100 file checks and evidence extraction, only in the isolated judge.

No submitted code is imported or executed. Semantic correctness is judged separately.
"""
import argparse
import csv
import hashlib
import json
import wave
from pathlib import Path


def text_of(path):
    suffix = path.suffix.lower()
    if suffix == '.pdf':
        import fitz
        with fitz.open(path) as doc:
            return '\n'.join(f'[page {i+1}]\n{page.get_text()}' for i, page in enumerate(doc))
    if suffix == '.docx':
        from docx import Document
        doc = Document(path)
        return '\n'.join([p.text for p in doc.paragraphs] +
                         [' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
    if suffix == '.pptx':
        from pptx import Presentation
        pres = Presentation(path)
        return '\n'.join(f'[slide {i+1}]\n' + '\n'.join(s.text for s in slide.shapes if s.has_text_frame) +
                         '\n[notes]\n' + (slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else '')
                         for i, slide in enumerate(pres.slides))
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            return '\n'.join(f'[{s.title}]\n' + '\n'.join(' | '.join(str(v) for v in row) for row in s.values) for s in book)
        finally:
            book.close()
    if suffix in {'.png', '.jpg', '.jpeg'}:
        from PIL import Image
        with Image.open(path) as im:
            return json.dumps({'width': im.width, 'height': im.height, 'format': im.format})
    if suffix == '.wav':
        with wave.open(str(path)) as wav:
            return json.dumps({'seconds': wav.getnframes()/wav.getframerate(), 'channels': wav.getnchannels()})
    if suffix in {'.db', '.sqlite'}:
        # Database answers are not executed here; tool traces and SQL output must support execution claims.
        return '[SQLite binary; inspect attached source if needed]'
    if path.stat().st_size > 2_000_000:
        return '[binary or oversized file; content omitted]'
    try:
        return path.read_text(encoding='utf-8')
    except UnicodeError:
        return '[binary file; content omitted]'


def extract(directory, *, output_errors=None):
    result = {}
    total = 0
    for p in sorted(Path(directory).rglob('*')):
        if p.is_file():
            if p.is_symlink():
                raise ValueError('symlink_in_evidence')
            try:
                value = text_of(p)
            except ImportError:
                raise
            except Exception as error:
                if output_errors is None:
                    raise
                output_errors.append({'path':p.relative_to(directory).as_posix(),
                                      'reason':type(error).__name__+':'+str(error)})
                value = '[unreadable submitted file]'
            total += len(value)
            # Refuse an incomplete semantic evidence set instead of silently truncating it.
            if total > 1_500_000:
                raise ValueError('business_evidence_too_large')
            result[p.relative_to(directory).as_posix()] = value
    return result


def verify(case, inputs, outputs, check):
    checks = []
    for name in case['expected_files']:
        p = outputs / name
        checks.append(check('非空交付：'+name, p.is_file() and p.stat().st_size > 0))
    n = int(case['source']['task'][1:])
    files = [p for p in outputs.rglob('*') if p.is_file()]
    if n == 63:
        import fitz
        with fitz.open(outputs/'summary.pdf') as pdf:
            checks.append(check('恰两页PDF', len(pdf) == 2, str(len(pdf))))
    if n == 66:
        from pptx import Presentation
        slides = Presentation(outputs/'slides.pptx').slides
        checks.append(check('10页可编辑文本幻灯片', len(slides) == 10 and all(
            any(s.has_text_frame and s.text.strip() for s in slide.shapes) for slide in slides)))
        checks.append(check('每页非空演讲备注', len(slides) == 10 and all(
            slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip() for slide in slides)))
    if n in {65, 100}:
        from docx import Document
        p = outputs/('manual.docx' if n == 65 else 'proposal.docx')
        doc = Document(p)
        checks.append(check('文档有可编辑正文', bool(''.join(x.text for x in doc.paragraphs).strip()) or bool(doc.tables)))
    if n == 100:
        from openpyxl import load_workbook
        book = load_workbook(outputs/'quotes.xlsx', data_only=False)
        try:
            checks.append(check('比较表含原生公式', any(c.data_type == 'f' for s in book for row in s for c in row)))
        finally:
            book.close()
    if n in {23, 28, 68, 71, 72}:
        names = [x for x in case['expected_files'] if x.endswith('.csv')]
        for name in names:
            with (outputs/name).open(encoding='utf-8-sig', newline='') as stream:
                rows = list(csv.reader(stream))
            checks.append(check('CSV有字段与记录：'+name, len(rows) >= 2 and bool(rows[0])))
            if n == 28:
                checks.append(check('宽转长字段和6条记录', rows[0] == ['product','month','amount'] and len(rows) == 7))
                try:
                    total = sum(float(row[2]) for row in rows[1:])
                except (ValueError, IndexError):
                    total = None
                checks.append(check('宽转长总额950', total == 950, str(total)))
    if n in {82, 83, 84, 88}:
        checks.append(check('实际交付Python源码', any(p.suffix == '.py' for p in files)))
    if n == 89:
        checks.append(check('实际交付修改后网页', any(p.suffix == '.html' for p in files)))
    if n == 90:
        checks.append(check('实际交付SQL', any(p.suffix == '.sql' for p in files)))
    if n in {91, 93}:
        folder = 'T091_downloads' if n == 91 else 'T093_rename'
        original = {p.relative_to(inputs/'samples'/folder).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
                    for p in (inputs/'samples'/folder).rglob('*') if p.is_file()}
        changed = {p.relative_to(outputs/folder).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
                   for p in (outputs/folder).rglob('*') if p.is_file()}
        from collections import Counter
        checks.append(check('合成文件完整且字节不变', Counter(original.values()) == Counter(changed.values()),
                            json.dumps({'before':original,'after':changed},ensure_ascii=False)))
        if n == 91:
            checks.append(check('实际移出原位置且分类到子目录', len(changed) == 3 and
                                all('/' in name for name in changed) and not (set(original) & set(changed))))
        else:
            renamed = '2026-10-02_项目A_会议.txt'
            old = 'A会议2026.10.02.txt'
            checks.append(check('完整日期文件实际重命名', changed.get(renamed) == original.get(old) and old not in changed))
            checks.append(check('日期或主题缺失文件保持原名', all(changed.get(k) == v for k,v in original.items() if k != old)))
    if n == 94:
        from PIL import Image
        images = [p for p in files if p.suffix.lower() in {'.jpg','.png','.jpeg'}]
        original = list((inputs/'samples/T094_images').glob('*'))
        actual = []
        for p in images:
            with Image.open(p) as im:
                actual.append(im.size)
        checks.append(check('三张实际缩放副本', len(original) == len(images) == 3 and
                            sorted(actual) == sorted([(1200,800),(800,1200),(1200,1200)]), str(actual)))
    return checks, {'method':'business_structural_v1', 'checks':checks}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--inputs', required=True)
    parser.add_argument('--outputs', required=True)
    parser.add_argument('--out', required=True)
    args = parser.parse_args()
    from .verify import preflight
    errors = []
    structural = preflight(Path(args.outputs))
    errors += [{'path':c['rule'],'reason':c.get('detail','invalid file')}
               for c in structural if not c['passed']]
    data = {'inputs':extract(Path(args.inputs)),
            'outputs':extract(Path(args.outputs),output_errors=errors), 'invalid_outputs':errors}
    Path(args.out).write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')


if __name__ == '__main__':
    main()
